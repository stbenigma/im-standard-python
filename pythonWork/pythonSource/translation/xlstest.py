import openpyxl
import re,os
import sys
import requests
import json

def translateString(pval,pfrom,pto):

    payload = {'auth_key': '9043d070-26fd-f874-6b80-37ddc4b6367c'
            , 'text': pval
            , 'target_lang': str.upper(pto)
            , 'source_lang': str.upper(pfrom)
            , 'preserve_formatting': 1}

    r = requests.post('https://api.deepl.com/v2/translate', params=payload)
    if (r.status_code == 200):
        rj = r.json()
        result = rj['translations'][0]['text']
    else:
        result = pval
    #fi
    return result
#translateString

def translateSheet (pfileName, pdestFileName
                    , pfromLang, ptoLang):
    wb = openpyxl.load_workbook(filename = pfileName)
    #print(wb.sheetnames)
    ws = wb.active
    #print(wb.sheetnames,ws)

    x=0
    for row in ws.values:
        x+=1
        y=0
        for val in row:
            y+=1
            if (val is not None):
                lregexpFromMarker = "\*{}\* ".format(str.upper(pfromLang))
                if (re.match(lregexpFromMarker, val)) :
                    newval = "**"+ translateString(pval=val[len(lregexpFromMarker)-2:],pfrom=pfromLang,pto=ptoLang)
                    #print(val,newval)
                    _ = ws.cell(column=y
                        , row=x
                         , value="{}".format(newval))
            #fi
        #for
    #for
    wb.save(filename=pdestFileName)
#translateSheet

def main():
    #print(    translateString("""""",'de','en')     )
    #return
    lfile =  sys.argv[1]
    lfromLang=  sys.argv[2]
    ltoLang =  sys.argv[3]
    lfileName, lfileExt = os.path.splitext(lfile)
    ldestFileName = lfileName + str.upper(ltoLang)+lfileExt

    translateSheet(pfileName=lfile, pdestFileName=ldestFileName
                   , pfromLang=lfromLang, ptoLang=ltoLang, )


if __name__ == '__main__':
    main()