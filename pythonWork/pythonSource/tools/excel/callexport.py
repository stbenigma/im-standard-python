# -*- coding: latin-1 -*-
import copy
import re
import datetime
import sys

from openpyxl import Workbook, styles
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter

from SSOT_db.IM_OBJECTS import Domain
from SSOT_infra import nvl
from tools.excel.datamapping import mappingexceldata as mapxls


fileCSV = None
EOL: str = '\n'
CSVSEP: str = ';'

greenfill = styles.PatternFill(patternType='solid',
                               fill_type='solid',
                               fgColor=styles.Color('80ED99'))
lightgrayfill = styles.PatternFill(patternType='solid',
                                   fill_type='solid',
                                   fgColor=styles.Color('EFEFEF'))
emptyfill = styles.PatternFill(fill_type=None)


def switchfill(cellfill: styles.PatternFill) -> styles.PatternFill:
    if cellfill == emptyfill:
        return lightgrayfill
    else:
        return emptyfill


thin_border = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))


def colnum_string(n):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string


def cell_string(n, m):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string + str(m)


def setcell(pws, pcolumn, prow, pvalue
            , phorizontal=None, pvertical=None
            , ptext_rotation=None, pwrap_text=None
            , ppattern=None):
    cell = pws.cell(column=pcolumn, row=prow, value=pvalue)
    if ppattern is not None:
        cell.fill = ppattern
        cell.border = thin_border
    if (phorizontal is not None or pvertical is not None
            or ptext_rotation is not None or pwrap_text is not None):
        cell.alignment = styles.Alignment(horizontal=phorizontal
                                          , vertical=pvertical
                                          , text_rotation=ptext_rotation
                                          , wrap_text=pwrap_text
                                          )
    """ Alignement horizontal  ?left?, ?centerContinuous?, ?center?, ?distributed?, ?fill?, ?justify?, ?right?, ?general?"""
    return


def writesheettabent(pwb: Workbook, pmodel, plang):
    ws = pwb.create_sheet(mapxls.WS_TAB2ENTMAP)
    rowidx, colidx = 1, 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Datamodel')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Table')
    colidx += 1

    cntcolidx, cntrowidx = colidx, rowidx + 1
    xcounts = {}
    setcell(pws=ws, prow=rowidx + 1, pcolumn=colidx, pvalue='Count', phorizontal='right')
    colidx += 1
    for idx, enti in enumerate(pmodel.getelements('entities').values()):
        xcounts[colidx + idx] = 0
        setcell(pws=ws, pcolumn=idx + colidx, prow=rowidx, pvalue=enti['name'][plang]
                , ptext_rotation=90)
    # for
    rowidx += 2
    for sval in pmodel.getelements("datamodels").values():
        for tkey in sval['tables+']:
            table = pmodel.getbyid(tkey)
            xcount, colidx = 0, 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=sval['name'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=table['name'])
            colidx += 2  # platz für counter
            for entiid in pmodel.getelements('entities').keys():
                if (entiid in table['entitiesmapped']):
                    val = table["CRUD"]
                    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=val,
                            phorizontal="center", pvertical="center",
                            ppattern=None if val is None else greenfill)
                    xcount += 1
                    xcounts[colidx] += 1
                # fi
                colidx += 1
            # for
            if (xcount > 0):
                setcell(pws=ws, prow=rowidx, pcolumn=cntcolidx
                        , pvalue=xcount, phorizontal='right')
            rowidx += 1
        # for
    # for
    for key, value in xcounts.items():
        if (value > 0):
            setcell(pws=ws, prow=cntrowidx, pcolumn=key, pvalue=value, phorizontal='right')
    rowidx += 1

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 7
    for idx in range(4, colidx): ws.column_dimensions[colnum_string(idx)].width = 5
    return


def writesheetentitab(pwb: Workbook, pmodel, plang):
    ws = pwb.create_sheet(mapxls.WS_ENT2TABMAP)

    rowidx, colidx = 1, 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Information Model')
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='Entity')
    colidx += 1

    cntcolidx, cntrowidx = colidx, rowidx + 2
    setcell(pws=ws, prow=rowidx + 2, pcolumn=colidx, pvalue='Count', phorizontal='right')
    xcounts = {}
    colidx += 1
    for sval in pmodel.getelements("datamodels").values():
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=sval['name'])
        for tkey in sval['tables+']:
            setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue=pmodel.getbyid(tkey)['name']
                    , ptext_rotation=90)
            xcounts[colidx] = 0
            colidx += 1
        # for
    # for
    rowidx += 3

    for entiid, enti in pmodel.getelements("entities").items():
        xcount, colidx = 0, 1
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=enti['name'][plang])
        colidx += 2  # platz fÃ¼r counter

        for sval in pmodel.getelements("datamodels").values():
            for tkey in sval['tables+']:
                if (entiid in pmodel.getbyid(tkey)['entitiesmapped']):
                    val = pmodel.getbyid(tkey)["CRUD"]
                    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=val,
                            ppattern=None if val is None else greenfill)
                    xcounts[colidx] += 1
                    xcount += 1
                # if
                colidx += 1
            # for
        # for
        if (xcount > 0):
            setcell(pws=ws, prow=rowidx, pcolumn=cntcolidx
                    , pvalue=xcount, phorizontal='right')
        rowidx += 1
    # for
    for key, value in xcounts.items():
        if (value > 0):
            setcell(pws=ws, prow=cntrowidx, pcolumn=key, pvalue=value, phorizontal='right')
    rowidx += 1

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 7
    for idx in range(3, colidx): ws.column_dimensions[colnum_string(idx)].width = 5
    return


# def writesheetcolattr(pwb: Workbook, pmodel, plang):
#     ws = pwb.create_sheet(mapxls.WS_COL2ATTRMAP)
#
#     rowidx, colidx = 1, 1
#     setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx, pvalue='Entity')
#     setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='Datamodel')
#     setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx + 1, pvalue='Table')
#     setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx + 1, pvalue='Column')
#     colidx += 3
#     for enti in pmodel.getelements("entities").values():
#         setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=enti['name'][plang])
#         for attr in enti['attributes']:
#             cell = setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue=pmodel.getbyid(attr)['name'][plang])
#             cell.alignment = Alignment(horizontal='general'
#                                        , vertical='bottom'
#                                        , text_rotation=90
#                                        # ,wrap_text = False
#                                        )
#             colidx += 1
#         # for
#     # for
#     rowidx += 2
#
#     for skey, sval in pmodel.getelements("datamodels").items():
#         for tkey in sval['tables+']:
#             colidx = 1
#             tabl = pmodel.getbyid(tkey)
#             for ckey in tabl['columns']:
#                 setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=sval['name'])
#                 colidx += 1
#                 setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=tabl['name'])
#                 colidx += 1
#                 setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=pmodel.getbyid(ckey)['name'])
#                 colidx += 1
#                 for enti in pmodel.getelements("entities").values():
#                     for attrid in enti['attributes']:
#                         attrcols = pmodel.getbyid(attrid)['columnsmapped']
#                         if (skey in attrcols.keys()) and (ckey in attrcols[skey]):
#                             setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='X')
#                         # if
#                         colidx += 1
#                     # for
#                 # for
#                 rowidx += 1
#             # for
#         # for
#     # for
#
#     ws.column_dimensions['A'].width = 35
#     ws.column_dimensions['B'].width = 35
#     ws.column_dimensions['C'].width = 35
#     for idx in range(4, colidx): ws.column_dimensions[colnum_string(idx)].width = 3
#     return


def valmiteol(pval, padd):
    eoladd = '' if (pval == '') else EOL
    return pval + eoladd + padd


def writedatmheader(pmodel, pws, pcolumn, prow, psystemids: list, pfirstpattern) -> (int, int):
    cellfill = pfirstpattern
    colidx, rowidx = pcolumn, prow
    for datm in psystemids:
        sval = pmodel.getbyid(datm)
        setcell(pws=pws, pcolumn=colidx, prow=rowidx, pvalue=sval['name'], ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 1, prow=rowidx, pvalue='', ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 2, prow=rowidx, pvalue='', ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 3, prow=rowidx, pvalue='', ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx, prow=rowidx + 1, pvalue="Table", ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 1, prow=rowidx + 1, pvalue="Column", ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 2, prow=rowidx + 1, pvalue="RW", ppattern=cellfill)
        setcell(pws=pws, pcolumn=colidx + 3, prow=rowidx + 1, pvalue="Ext-ID", ppattern=cellfill)
        colidx += 4
        cellfill = switchfill(cellfill)
    # for
    return (colidx, rowidx)

def getcolmappings(pmodel, pdatmid, pattrid) -> list:
    valuetab, valuecol, valuerw, valueid = '', '', '', ''
    for ckey, colu in pmodel.getelements('columns').items():
        if (colu["datamodel-id+"] == pdatmid) and \
                (pattrid in [c[0] for c in colu['attributesmapped']]):
            valuetab = valmiteol(pval=valuetab, padd=colu['table-name+'])
            valuecol = valmiteol(pval=valuecol, padd=colu['name'])
            valuerw = valmiteol(pval=valuerw, padd=nvl(colu['R/W']))
            valueid = valmiteol(pval=valueid, padd=nvl(colu['datamodel_col_id']))
        # if
    # for
    return [valuetab, valuecol, valuerw, valueid]


def writesheetattrcol(pwb: Workbook, pmodel, psysids, plang):
    ws = pwb.create_sheet(mapxls.WS_ATTR2COLMAP)

    rowidx, colidx = 1, 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Information Model')
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='Entity')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx + 1, pvalue='Attribute')
    colidx += 2
    colidx, rowidx = writedatmheader(pmodel=pmodel, pws=ws, pcolumn=colidx, prow=rowidx,
                                     psystemids=psysids, pfirstpattern=lightgrayfill
                                     )
    rowidx += 2

    attrsort = [[key, pmodel.getbyid(val["entity"])["name"][plang], val['name'][plang]] for key, val in
                pmodel.getelements("attributes").items()]
    attrsort = sorted(attrsort, key=lambda x: x[1] + "-" + x[2])
    for attr in attrsort:
        attrid, entiname, attrname = attr[0], attr[1], attr[2]
        colidx = 1
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=entiname, pwrap_text=True)
        colidx += 1
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=attrname, pwrap_text=True)
        colidx += 1
        cellfill = emptyfill
        for skey in pmodel.getelements("datamodels").keys():
            cellfill = switchfill(cellfill)
            valuetab, valuecol, valuerw, valueid = getcolmappings(pmodel=pmodel, pdatmid=skey, pattrid=attrid)
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuetab
                    , pwrap_text=True, ppattern=cellfill)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuecol
                    , pwrap_text=True, ppattern=cellfill)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuerw
                    , pwrap_text=True, ppattern=cellfill)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valueid
                    , pwrap_text=True, ppattern=cellfill)
            colidx += 1
            headcell = ws.cell(column=colidx, row=3)
            headcell.alignment = styles.Alignment(text_rotation=0, wrap_text=True)
        # for
        rowidx += 1
    # for

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    for idx in range(3, colidx):
        colwidth = {0: 35, 1: 5, 2: 15, 3: 35}
        """3 4 .5 6 7   8 .9 10 11"""
        ws.column_dimensions[colnum_string(idx)].width = colwidth[idx % 4]
    # for
    return


def writesheetdatamodel(pwb, pdatmid, pmodel, plang, pothersysids):
    datm = pmodel.getbyid(pdatmid)
    ws = pwb.create_sheet(datm["name"][:31]) #max length for tab names

    rowidx, colidx = 1, 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=datm['name'])
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='tableName')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='tableCRUD')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='columnName')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='columnRW')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='column-ID')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='domain')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='dataType')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='mand.')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='default value')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='descr')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='rules')
    colidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Information Model', ppattern=lightgrayfill)
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue='', ppattern=lightgrayfill)
    setcell(pws=ws, pcolumn=colidx, prow=rowidx + 1, pvalue='entityName', ppattern=lightgrayfill)
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx + 1, pvalue='attrName', ppattern=lightgrayfill)
    colidx += 2

    colidx, rowidx = writedatmheader(pws=ws, pmodel=pmodel, pcolumn=colidx, prow=rowidx,
                                     psystemids=pothersysids, pfirstpattern=emptyfill)
    rowidx += 2

    for tabid in datm['tables+']:
        tabl = pmodel.getbyid(tabid)
        firstrowidx = rowidx
        for colid in tabl['columns+']:
            colu = pmodel.getbyid(colid)
            valueenti, valueattr = '', ''
            attrids = []
            for attrmap in colu['attributesmapped']:
                attr = pmodel.getbyid(attrmap[0])
                attrids.append(attrmap)
                assert attr is not None, f"Missing attribute '{attrmap[0]}' from column '{colid}'"
                enti = pmodel.getbyid(attr['entity'])
                valueenti = valmiteol(pval=valueenti, padd=enti['name'][plang] + (
                    '' if attrmap[1] is None else f" ({pmodel.getbyid(attrmap[1])['name'][plang]})"))
                valueattr = valmiteol(pval=valueattr, padd=attr['name'][plang])
            # for
            colidx = 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=tabl['name'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=tabl['CRUD'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=colu['name'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=colu['R/W'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=colu['datamodel_col_id'])
            colidx += 1
            doma = pmodel.getbyid(colu['domain'])
            setcell(pws=ws, pcolumn=colidx, prow=rowidx,
                    pvalue=doma['name'][plang] if doma['origin'] == Domain.DOMAIN else None)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=colu['datatype'])
            colidx += 1
            # setcell(pws=ws,pcolumn=colidx,prow=rowidx,pvalue=colu['mandatory'])
            colidx += 1
            # setcell(pws=ws,pcolumn=colidx,prow=rowidx,pvalue=colu['default Value'])
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=colu['descr'], pwrap_text=True)
            colidx += 1
            # setcell(pws=ws,pcolumn=colidx, prow=rowidx, pvalue=colu['rules'],pwrap_text=True)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valueenti, pwrap_text=True, ppattern=lightgrayfill)
            colidx += 1
            setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valueattr, pwrap_text=True, ppattern=lightgrayfill)
            colidx += 1
            cellfill = lightgrayfill
            for datmid in pothersysids:  # for all other datamodels
                cellfill = switchfill(cellfill)
                valuetab, valuecol, valuerw, valueid = '', '', '', ''
                for attrmap in attrids:
                    valuetab1, valuecol1, valuerw1, valueid1 = getcolmappings(pmodel=pmodel, pdatmid=datmid,
                                                                              pattrid=attrmap[0])
                    if valuetab1 != '':
                        valuetab = valmiteol(valuetab, valuetab1)
                        valuecol = valmiteol(valuecol, valuecol1)
                        valuerw = valmiteol(valuerw, valuerw1)
                        valueid = valmiteol(valueid, valueid1)
                setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuetab
                        , pwrap_text=True, ppattern=cellfill)
                colidx += 1
                setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuecol
                        , pwrap_text=True, ppattern=cellfill)
                colidx += 1
                setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valuerw
                        , pwrap_text=True, ppattern=cellfill)
                colidx += 1
                setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=valueid
                        , pwrap_text=True, ppattern=cellfill)
                colidx += 1
            # for

            rowidx += 1
        # for
        if firstrowidx == rowidx:
            """no columns written write the entity-Mappings if present"""
            valueenti = ''
            for entiid in tabl['entitiesmapped']:
                enti = pmodel.getbyid(entiid)
                valueenti = valmiteol(pval=valueenti, padd=enti['name'][plang])
            # for
            setcell(pws=ws, pcolumn=1, prow=rowidx, pvalue=tabl['name'])
            setcell(pws=ws, pcolumn=2, prow=rowidx, pvalue=tabl['CRUD'])
            setcell(pws=ws, pcolumn=12, prow=rowidx, pvalue=valueenti, pwrap_text=True, ppattern=lightgrayfill)
            setcell(pws=ws, pcolumn=13, prow=rowidx, pvalue='', pwrap_text=True, ppattern=lightgrayfill)
            cellfill = lightgrayfill
            colidx = 14
            for idx in range(len(pothersysids)):  # for all other datamodels
                cellfill = switchfill(cellfill)
                for f in range(4):
                    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='', pwrap_text=False, ppattern=cellfill)
                    colidx += 1

            rowidx += 1
        # if
    # for
    colsizes1 = [35, 10, 40, 10, 20,
                 20, 20, 10, 20, 35,
                 35, 35, 35]
    colsizes2 = [35, 35, 6, 25]  # for datamodel-attributes
    colidx = 1
    for width in colsizes1:
        ws.column_dimensions[get_column_letter(colidx)].width = width
        colidx += 1
    for datm in pothersysids:
        for width in colsizes2:
            ws.column_dimensions[get_column_letter(colidx)].width = width
            colidx += 1
    return


# def writexls(pfilename: str, pmodel, plang):
#     wb = Workbook()
#     writesheettabent(pwb=wb, pmodel=pmodel, plang=plang)
#     writesheetentitab(pwb=wb, pmodel=pmodel, plang=plang)
#     writesheetattrcol(pwb=wb, pmodel=pmodel, plang=plang, psysids=psysids)
#     #    writesheetcolattr(pwb=wb)
#     wb.remove(wb.worksheets[0])
#     wb.save(filename=pfilename)
#     return


def writeoverview(pwb, pmodel, plang):
    systems = pmodel.getelements('DATM')
    entis = pmodel.getelements('ENTI')
    relas = pmodel.getelements('RELA')
    attrs = pmodel.getelements('ATTR')
    tabs = pmodel.getelements('TABL')
    cols = pmodel.getelements('COLU')
    ws = pwb.create_sheet(mapxls.WS_OVERVIEW)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    rowidx, colidx = 1, 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='System')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue='Tables')
    setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx, pvalue='Enti/Rela mapped')
    setcell(pws=ws, pcolumn=colidx + 3, prow=rowidx, pvalue='Columns')
    setcell(pws=ws, pcolumn=colidx + 4, prow=rowidx, pvalue='Attributes mapped')
    for syskey, sys in systems.items():
        rowidx += 1
        c = ws.cell(column=1, row=2)
        c.alignment = styles.Alignment(horizontal='general'
                                       , vertical='bottom'
                                       , text_rotation=0,
                                       wrap_text=False,
                                       shrink_to_fit=False,
                                       indent=0)
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=sys['name'])
        setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=len(sys["tables+"]))
        setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx
                , pvalue=sum([len(t["entitiesmapped"]) + len(t["relationsmapped"]) for t in tabs.values() if
                              t["datamodel-id"] == syskey]))
        setcell(pws=ws, pcolumn=colidx + 3, prow=rowidx
                , pvalue=len([c["name"] for c in cols.values() if c["datamodel-id+"] == syskey]))
        setcell(pws=ws, pcolumn=colidx + 4, prow=rowidx
                , pvalue=sum([len(c["attributesmapped"]) for c in cols.values() if c["datamodel-id+"] == syskey]))
    # for
    rowidx += 2
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Information Model')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue='Entities/Relations')
    setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx, pvalue='Tables mapped')
    setcell(pws=ws, pcolumn=colidx + 3, prow=rowidx, pvalue='Attrbutes')
    setcell(pws=ws, pcolumn=colidx + 4, prow=rowidx, pvalue='Columns-Mapped')
    rowidx += 1
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=len(entis))
    setcell(pws=ws, pcolumn=colidx + 2, prow=rowidx, pvalue=sum([len(e["tablesmapped+"]) for e in entis.values()])
                                                            + sum([len(e["tablesmapped+"]) for e in relas.values()]))
    setcell(pws=ws, pcolumn=colidx + 3, prow=rowidx, pvalue=len(attrs))
    setcell(pws=ws, pcolumn=colidx + 4, prow=rowidx, pvalue=sum([len(a["columnsmapped+"]) for a in attrs.values()]))

    rowidx += 3
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Model')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=str(pmodel.modelname()))
    rowidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='JSON file')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=str(pmodel.jsfile))
    rowidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='JSON Version')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=pmodel.getjsversion().base_version)
    rowidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='created')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=str(datetime.datetime.now()))
    rowidx += 1
    setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue='Language')
    setcell(pws=ws, pcolumn=colidx + 1, prow=rowidx, pvalue=plang)

    return


def writedatmxls(pfilename: str, pmodel, plang):
    #    writesheetcolattr(pwb=wb)
    wb = Workbook()
    orderedsysids = [k for k, v in sorted(pmodel.getelements("datamodels").items(),
                                          key=lambda item: item[1]["name"])]
    writeoverview(pwb=wb, pmodel=pmodel, plang=plang)
    writesheettabent(pwb=wb, pmodel=pmodel, plang=plang)
    writesheetentitab(pwb=wb, pmodel=pmodel, plang=plang)
    writesheetattrcol(pwb=wb, pmodel=pmodel, plang=plang, psysids=orderedsysids)

    for datmid in orderedsysids:
        othersys = copy.copy(orderedsysids)
        othersys.remove(datmid)
        writesheetdatamodel(pwb=wb, pdatmid=datmid, pmodel=pmodel, plang=plang,
                            pothersysids=othersys)
    wb.remove(wb.worksheets[0])
    wb.save(filename=pfilename)
    return


# def istintabentimap(tabid, entiid):
#     return (tabid in tabentimap) and (entiid in tabentimap[tabid])


# def istincolattrmap(colid, attrid):
#     return (colid in colattrmap) and (attrid in colattrmap[colid])

def stripeol(str):
    retval = re.sub("\n+", "\n", str)
    retval = retval.strip(EOL)
    return retval



def main():
    from tools.excel import listmapping
    #redirect to general excel call
    sys.argv.insert(1,"--exporttype")
    sys.argv.insert(2,"datamapping")
    listmapping.main()

if __name__ == '__main__':
    main()
