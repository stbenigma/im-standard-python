import os.path
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from SSOT_db.IM_JSON import JSModel
from SSOT_infra.tests import integration as tb
from tools.excel.datamapping import listmapping


class ListMapping(unittest.TestCase):

    def setUp(self) -> None:
        self.localrun = os.path.isdir(Path.home()/"Downloads")
        self.crm = tb.ModelHelper(tb.CRMTEST)
        self.crm.initDB(palways=False)
        return

    def test_createAllMapping(self):
        with tempfile.TemporaryDirectory() as tempdir:
            testfile = Path(tempdir + "/testfile.json")
            shutil.copy(self.crm.jsonfile, testfile)
            mapfile = listmapping.createAllMapping(pjsonfile=testfile,
                                                   pdestination=testfile.with_suffix('.xlsx'))
            wb = load_workbook(filename=mapfile)

            # DEBUG
            if self.localrun:
                shutil.copy(mapfile,Path.home()/"Downloads")
                print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

            self.assertEqual("Overview", wb.sheetnames[0])
            wssf = wb.get_sheet_by_name("CRM-Salesforce")
            self.assertEqual("person", wssf.cell(153, 1).value)
            self.assertEqual("Geschlecht", wssf.cell(153, 3).value)
            self.assertEqual("natürliche Person", wssf.cell(153, 12).value)
            self.assertEqual("Geschlecht", wssf.cell(153, 13).value)
            self.assertEqual("*natural_person", wssf.cell(153, 18).value)
            self.assertEqual("*Gender", wssf.cell(153, 19).value)

            locmodel = JSModel.readfromfile(self.crm.jsonfile)
            tb.crmmapentityhack(locmodel)
            mapfile = Path(tempdir + '/map.xlsx')
            listmapping.writedatmxls(pfilename=str(mapfile), pmodel=locmodel,
                                     plang=locmodel.jsmodel["model"]["language"])
            wb = load_workbook(filename=mapfile)
            if self.localrun:
                shutil.copy(mapfile,Path.home()/"Downloads")
                print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

            wssf = wb.get_sheet_by_name('Overview')
            self.assertEqual(locmodel.jsmodel["model"]["language"], wssf.cell(16, 2).value)

            wssf = wb.get_sheet_by_name("CRM-Salesforce")
            self.assertEqual("GeoInformation", wssf.cell(43, 3).value)
            self.assertEqual("*col1tomany", wssf.cell(43, 27).value)

            wssf = wb.get_sheet_by_name("TestMapping")
            self.assertEqual("col1tomany", wssf.cell(4, 3).value)
            self.assertTrue(wssf.cell(4, 12).value.startswith("Geografische Einheit (Administrativgebiet)"))
            self.assertTrue(wssf.cell(4, 13).value.startswith("GeoInformation"))

            listmapping.writedatmxls(pfilename=str(mapfile), pmodel=locmodel,
                                     plang="en")
            wb = load_workbook(filename=mapfile)

            wssf = wb.get_sheet_by_name('Overview')
            self.assertEqual("en", wssf.cell(16, 2).value)

            wssf = wb.get_sheet_by_name("TestMapping")
            self.assertTrue( wssf.cell(4, 12).value.endswith("Geografical Unit (Administrative Territory)"))
            if self.localrun:
                shutil.copy(mapfile,Path.home()/"Downloads")
                print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

        return

    def test_excel_names(self):
        with tempfile.TemporaryDirectory() as tempdir:
            locmodel = JSModel.readfromfile(self.crm.jsonfile)

            syst = locmodel.getelements("datamodels")
            locmodel.getbyid(list(syst.keys())[1])[
                "name"] += "123456789012345678901234567890"  # make too long system name
            mapfile = Path(tempdir + '/map.xlsx')

            listmapping.writedatmxls(pfilename=str(mapfile), pmodel=locmodel,
                                     plang=locmodel.jsmodel["model"]["language"])
        return

    def test_specialjson(self):
        # return  # for local customer-jsonss
        self.skipTest("for special test going over several tools")

        with tempfile.TemporaryDirectory() as tempdir:
            if self.localrun:
                testfile = tempdir + "/testfile.json"
                shutil.copy('/Users/stb/Downloads/spoddataspotjson.json', testfile)
                mapfile = listmapping.createAllMapping(pjsonfile=testfile,pdestination=tempdir + "/testfile.xlsx")
                self.assertTrue(os.path.exists((tempdir + "/testfile.xlsx")))
                shutil.copy(tempdir + "/testfile.xlsx","/Users/stb/Downloads/spoddataspotjson.xlsx")
            return

    def test_call(self):
        sys.argv = [
            listmapping.__file__,
            '--language', 'en',
            str(self.crm.jsonfile)
        ]
        listmapping.main()
        sys.argv = [
            listmapping.__file__,
            str(self.crm.jsonfile)
        ]
        with tempfile.TemporaryDirectory() as tempdir:
            testfile = tempdir + "/maptestexcel.xlsx"
            listmapping.main()
            sys.argv = [
                listmapping.__file__,
                '--destination', testfile,
                str(self.crm.jsonfile)
            ]
            listmapping.main()
            self.assertTrue(os.path.isfile(testfile))
            if self.localrun:
                shutil.copy(testfile,Path.home()/"Downloads")

            #test diagrams and stat filters
            sys.argv = [
                listmapping.__file__,
                '-s','GTOP',
                '--diagrams=DUMMY',
                '--destination',testfile,
                str(self.crm.jsonfile)
            ]
            listmapping.main()
            self.assertTrue(os.path.isfile(testfile))
            if self.localrun:
                shutil.copy(testfile,Path.home()/"Downloads"/"filtered.xlsx")

        return


if __name__ == '__main__':
    unittest.main()
