import logging
import os.path
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

from SSOT_db.IM_JSON import JSModel
from SSOT_infra.tests import integration as tb
from tools.excel.datamapping import listmapping, importMAP
from tools.excel.datamapping import mappingexceldata as mapxls
from SSOT_infra import nvl



class ListImportMapping(unittest.TestCase):
    def setUp(self) -> None:
        root = logging.getLogger()
        root.setLevel(logging.INFO)
        handler = logging.StreamHandler(sys.stdout)
        handler.setLevel(logging.INFO)
        root.addHandler(handler)

        self.crm = tb.ModelHelper(tb.CRMTEST)

        self.crmjson = JSModel.readfromfile(self.crm.jsonfile)
        self.crmxls = self.crm.jsonfile.with_name(self.crm.jsonfile.stem + "_datamodels").with_suffix(".xlsx")
        listmapping.createAllMapping(pjsonfile=self.crm.jsonfile,
                                     pdestination=self.crmxls)
        self.assertTrue(os.path.exists(self.crmxls))
        return

    def test_findrowcol(self):
        wb: openpyxl.workbook.Workbook = openpyxl.load_workbook(self.crmxls)
        ws: openpyxl.openpyxl.Worksheet = wb[mapxls.WS_OVERVIEW]
        r = importMAP.findrow(ws, {1: "JSON file"})
        self.assertEqual(ws.cell(13, 2), r[2 - 1])
        c = importMAP.findcol(ws, {1: "Columns", 8: "Attributes"})
        self.assertEqual(ws.cell(2, 4), c[2 - 1])
        return

    def test_importmain(self):
        with self.assertRaises(AssertionError) as ae:
            importMAP.importmapexcel(pexcelfile=None)
        with self.assertRaises(AssertionError) as ae:
            importMAP.importmapexcel(pexcelfile=self.crmxls, pjsonfile="gugus")

        # nimm json-file aus dem Excel
        self.assertEqual(0, importMAP.importmapexcel(pexcelfile=self.crmxls))

        with tempfile.TemporaryDirectory() as tempdir:
            # preapre hack for missing subentity mappings
            testexcel = Path(tempdir + "/maptestexcel.xlsx")
            testjson = Path(tempdir + "/testjson.json")
            testdb = Path(tempdir + "/testmodel.db")

            locmodel = JSModel.readfromfile(self.crm.jsonfile)
            tb.crmmapentityhack(locmodel)
            locmodel.write_json(testjson)
            listmapping.writedatmxls(pfilename=testexcel, pmodel=locmodel,
                                     plang=locmodel.jsmodel["model"]["language"])

            #shutil.copy(testjson, Path.home() / "Downloads")
            #shutil.copy(testexcel,Path.home()/"Downloads")
            #print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")
            # nimm mein lokales json-file
            changes = importMAP.importmapexcel(pexcelfile=testexcel, pjsonfile=testjson)
            #shutil.copy(testjson, Path.home() / "Downloads/new.json")
            self.assertEqual(0, changes)

            # check  before the local change
            checkmodel = JSModel.readfromfile(testjson)
            tab = checkmodel.getbyfield(ptype="tables", pvalue="TabRela")
            enti = checkmodel.getbyfield(ptype="entities", pvalue="Personenrolle", plang="de")
            self.assertNotIn(enti[0][0], tab[0][1]["entitiesmapped"])

            wb = openpyxl.load_workbook(testexcel)
            ws = wb["TestMapping"]
            ws.cell(6, 12).value = nvl(ws.cell(6, 12).value) + "\nLand"
            ws.cell(10, 12).value = "Personenrolle"
            # absichtlich Zeile leer gelassen.
            ws.cell(12, 1).value = "ColAttr"
            ws.cell(12, 3).value = "Col1to1"
            ws.cell(12, 12).value = "Personenrolle"
            ws.cell(12, 13).value = "Dauer"
            wb.save(testexcel)
            shutil.copy(self.crm.dbfile, testdb)
            changes = importMAP.importmapexcel(pexcelfile=testexcel, pjsonfile=testjson, pdbfile=testdb)
            #shutil.copy(testexcel, os.path.expanduser('~') + "/Downloads/")
            #shutil.copy(testjson, os.path.expanduser('~') + "/Downloads/")
            #shutil.copy(testdb, os.path.expanduser('~') + "/Downloads/")
            self.assertEqual(4, changes)
            checkmodel = JSModel.readfromfile(testjson)
            tab = checkmodel.getbyfield(ptype="tables", pvalue="ColAttr")
            col = checkmodel.getbyfield(ptype="columns", pvalue="Col1to1")
            enti = checkmodel.getbyfield(ptype="entities", pvalue="natürliche Person", plang="de")
            attr = [a for a in enti[0][1]["attributes+"] if checkmodel.getbyid(a)["name"]['de'] == "Vorname"]
            self.assertIn(enti[0][0], tab[0][1]["entitiesmapped"])
            self.assertIn(attr[0], [c[0] for c in col[0][1]["attributesmapped"]])

        return

    def test_call(self):
        from SSOT_db import createJSON
        sys.argv = [
            importMAP.__file__,
            str(self.crmxls)
        ]
        importMAP.main()

        with tempfile.TemporaryDirectory() as tempdir:
            testexcel = Path(tempdir + "/maptestexcel.xlsx")
            testjson = Path(tempdir + "/testjson.json")
            testdb = Path(tempdir + "/testmodel.db")
            shutil.copy(self.crm.dbfile, testdb)
            createJSON.createJSON(pdbfilepath=testdb,
                                  pmodelname=self.crm.modelname,
                                  pjsfilepath=tempdir,
                                  pjsfilename="testjson.json")
            listmapping.createAllMapping(pjsonfile=testjson,
                                         pdestination=testexcel)

            sys.argv = [
                importMAP.__file__,
                "--dryrun",
                "--dbfile", str(testdb),
                "-j", str(testjson),
                str(testexcel)
            ]
            importMAP.main()

            wb = openpyxl.load_workbook(testexcel)
            ws = wb["TestMapping"]
            ws.cell(6, 12).value = nvl(ws.cell(6, 12).value) + "\nLand"
            ws.cell(10, 12).value = "Personenrolle"
            wb.save(testexcel)

            sys.argv = [
                importMAP.__file__,
                "--dryrun",
                "--dbfile", str(testdb),
                "-j", str(testjson),
                str(testexcel)
            ]
            importMAP.main()

            return

    def test_regexpmultiline(self):
        self.assertListEqual([], importMAP.getmultilines(pvalue=""))
        self.assertListEqual([["A", '']], importMAP.getmultilines(pvalue="A"))
        self.assertListEqual([["A", ''], ['B', '']], importMAP.getmultilines(pvalue="A\n  B  "))
        self.assertListEqual([["A", 'x'], ['B', '']], importMAP.getmultilines(pvalue="A(x)\n  B  ", pidx=None))
        self.assertListEqual([["A", 'x'], ['B', '']], importMAP.getmultilines(pvalue="A (x)\n  B\n  "))
        self.assertListEqual([['B', '']], importMAP.getmultilines(pvalue="(x)\n  B\n  ", pidx=None))

        self.assertListEqual([], importMAP.getmultilines(pvalue="", pidx=0))
        self.assertListEqual(["", ""], importMAP.getmultilines(pvalue="", pidx=2))
        self.assertListEqual(["A"], importMAP.getmultilines(pvalue="A", pidx=1))
        self.assertListEqual(["A", "", ""], importMAP.getmultilines(pvalue="A", pidx=3))
        self.assertListEqual(["A", 'B'], importMAP.getmultilines(pvalue="A\nB", pidx=0))
        self.assertListEqual(["A", 'B'], importMAP.getmultilines(pvalue="  \tA\nB", pidx=1))
        self.assertListEqual(["A", 'B'], importMAP.getmultilines(pvalue="A\nB\t  ", pidx=2))
        self.assertListEqual(["A", 'B', ''], importMAP.getmultilines(pvalue="A\n B ", pidx=3))
        self.assertListEqual(["A", 'B', 'C'], importMAP.getmultilines(pvalue="A\n  B  \nC", pidx=1))

        # importMAP.getmultilines(pvalue=, pidx: int = None)->list:
        return
