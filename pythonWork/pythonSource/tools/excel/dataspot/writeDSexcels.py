import os
import sys

from openpyxl import Workbook, styles
from openpyxl.styles.borders import Border, Side

from SSOT_db.IM_JSON import JSModel, readjsonfile

DEFAULTEXCELSTRUCT = os.path.dirname(os.path.realpath(__file__)) + "/fyayc_excelstruct.json"


def createexcel(pfilename, pmodel, plang, pfillfunc, pstruct):
    wb = Workbook()
    pfillfunc(pwb=wb, pmodel=pmodel, plang=plang, pstruct=pstruct)
    wb.remove(wb.worksheets[0])
    wb.save(filename=pfilename)
    print(f"\nExcel written {pfilename}")
    return


thin_border = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))


def colnum_string(n):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string


def cell_string(n, m):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string + str(m)


def setcell(pws, pcolumn, prow, pvalue
            , phorizontal=None, pvertical=None
            , ptext_rotation=None, pwrap_text=None
            , ppattern=None):
    cell = pws.cell(column=pcolumn, row=prow, value=pvalue)
    if ppattern is not None:
        cell.fill = ppattern
        cell.border = thin_border
    if (phorizontal is not None or pvertical is not None
            or ptext_rotation is not None or pwrap_text is not None):
        cell.alignment = styles.Alignment(horizontal=phorizontal
                                          , vertical=pvertical
                                          , text_rotation=ptext_rotation
                                          , wrap_text=pwrap_text
                                          )
    """ Alignement horizontal  ?left?, ?centerContinuous?, ?center?, ?distributed?, ?fill?, ?justify?, ?right?, ?general?"""
    return


def initsheet(pwb,pstruct):
    ws = pwb.create_sheet(pstruct["tab"])
    rowidx, colidx = 2, 1
    for head in pstruct["columns"]:
        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=head)
        colidx += 1
    return ws

def fillsheet(pwb,pstruct,data):
    ws=initsheet(pwb=pwb, pstruct=pstruct)
    rowidx= 3
    for row in data:
        colidx=1
        for val in row:
            setcell(pws=ws,pcolumn=colidx,prow=rowidx,pvalue=val)
            colidx+=1
        rowidx+=1

def getvalue(pdict,pkey,pfield,plang):
    if pkey is None :
        return None
    else:
        entry =pdict[pkey][pfield]
        if type(entry) == dict and plang in entry:
            return entry[plang]
        else:
            return entry

def writeBusinessModel(pdestdirec, pmodel, plang, pstruct):
    def dm2excel(pwb: Workbook, pmodel: JSModel, plang, pstruct):
        catg = pmodel.getelements("categories")
        fillsheet(pwb=pwb,pstruct=pstruct["Collections"],
                    data=[[val["name"],"Geschäftsobjekt-Kategorie",
                           '','','',key,
                           '','','','',
                           "In Arbeit",val["uc"],val["dc"][:19],
                           '',
                           "#"+val["ui"]["color"]] for key,val in catg.items()
                          ]
                  )
        entis = pmodel.getelements("entities")
        fillsheet(pwb=pwb,pstruct=pstruct["BusinessObjects"],
                  data=[[val["name"][plang],
                         getvalue(pdict=entis, pkey=val["supertypeentity"], pfield="name", plang=plang),
                         val["tooltip"][plang],
                         val["descr"][plang],
                         ",".join(s[plang] for s in val["synonyms"]),
                         ",".join(s[plang] for s in val["examples"]),
                        key,
                         catg[val["category"]]["name"],'','','',
                          "In Arbeit",val["uc"], val["dc"][:19],
                         ''
                         ] for key,val in entis.items()])
        entis = pmodel.getelements("entities")
        fillsheet(pwb=pwb,pstruct=pstruct["BusinessObjects"],
                  data=[[val["name"][plang],
                         getvalue(pdict=entis, pkey=val["supertypeentity"], pfield="name", plang=plang),
                         val["tooltip"][plang],
                         val["descr"][plang],
                         ",".join(s[plang] for s in val["synonyms"]),
                         ",".join(s[plang] for s in val["examples"]),
                        key,
                         catg[val["category"]]["name"],'','','',
                          "In Arbeit",val["uc"], val["dc"][:19],
                         ''
                         ] for key,val in entis.items()])

        initsheet(pwb=pwb,pstruct=pstruct["BusinessAttributes"])
        initsheet(pwb=pwb,pstruct=pstruct["Relationships"])
        initsheet(pwb=pwb,pstruct=pstruct["Transformations"])
        initsheet(pwb=pwb,pstruct=pstruct["TransformationRules"])
        return

    createexcel(pfilename=pdestdirec + "/" + pmodel.modelname() + "_BM.xlsx",
                pmodel=pmodel, plang=plang, pfillfunc=dm2excel, pstruct=pstruct)
    return


def writeDataModels(pdestdirec, pmodel, plang, pstruct):
    return


def writeDataTypes(pdestdirec, pmodel, plang, pstruct):
    return


def writeReferences(pdestdirec, pmodel, plang, pstruct):
    return


def writeTransformations(pdestdirec, pmodel, plang, pstruct):
    return


def writeDSfiles(pdestdirec, pmodel: JSModel, plang, pexcellayout, pexcellang="de"):
    excelstruct = readjsonfile(pfilename=DEFAULTEXCELSTRUCT if pexcellayout is None else pexcellayout)
    writeBusinessModel(pdestdirec=pdestdirec, pmodel=pmodel, plang=plang,
                       pstruct=excelstruct["BusinessModel"][pexcellang])
    writeDataModels(pdestdirec=pdestdirec, pmodel=pmodel, plang=plang,
                    pstruct=excelstruct["BusinessModel"][pexcellang])
    writeDataTypes(pdestdirec=pdestdirec, pmodel=pmodel, plang=plang,
                   pstruct=excelstruct["BusinessModel"][pexcellang])
    writeReferences(pdestdirec=pdestdirec, pmodel=pmodel, plang=plang,
                    pstruct=excelstruct["BusinessModel"][pexcellang])
    writeTransformations(pdestdirec=pdestdirec, pmodel=pmodel, plang=plang,
                         pstruct=excelstruct["BusinessModel"][pexcellang])
    return


def main():
    from tools.excel import callexport
    # redirect to general excel call
    sys.argv.insert(1, "--exporttype")
    sys.argv.insert(2, "dataspot")
    callexport.main()


if __name__ == '__main__':
    main()
