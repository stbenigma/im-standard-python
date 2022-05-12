import os.path
import sys
from openpyxl import Workbook
from SSOT_db.IM_JSON import JSModel

def excelfilename(pfilename):
    return pfilename + '.xlsx'

val2str = lambda v: '' if v is None else str(v) if type(v) in (int,float) else v
def writesheets(pwb):
    tables = gettablelist()
    tables = [tab.upper() for tab in tables]
    tables = sorted(tables)
    ws = pwb.create_sheet("Overview")
    rowidx, colidx = 1, 1
    ws.cell(column=colidx, row=rowidx, value='Table')
    ws.cell(column=colidx+1, row=rowidx, value='rowcount')
    for tab in tables:
        rowidx += 1
        ws.cell(column=colidx, row=rowidx, value=tab)
        ws.cell(column=colidx+1, row=rowidx, value=dbDML.getrowcount(ptablename=tab))
    #for
    for tab in tables:
        ws = pwb.create_sheet(tab)
        cols = dbDDL.getcolums(ptablename=tab)
        rowidx, colidx = 1, 2
        #headers, starting at 2, leaving room for CRUD column for importing
        for col in cols.keys():
            ws.cell(column=colidx, row=rowidx, value=col)
            colidx += 1
        #for
        data = dbDML.select(psql=f"select * from {tab}")
        for row in data:
            rowidx += 1
            colidx = 2
            for v in row:
                ws.cell(column=colidx, row=rowidx, value=val2str(v))
                colidx +=1
            #for
        # for
    return

def createExcel(pfilename: str):
    wb = Workbook()
    writesheets(pwb=wb)
    wb.remove(wb.worksheets[0])
    wb.save(filename=pfilename)
    return

def createlangexcels(pjsonfilepath,pdest=None):
    asserts os.path.exists(pjsonfilepath)
    print (os.path.dirname(pjsonfilepath))
    jsmodel = JSModel.readfromfile(pfilename=pjsonfilepath)
    return

if __name__ == '__main__':
    main(pjsonfilepath=sys.argv[1],pdest=None if len(sys.argv)<3 else sys.argv[2])
