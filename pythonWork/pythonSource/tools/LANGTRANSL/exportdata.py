import os.path
import sys

from openpyxl import Workbook, styles
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from tools.LANGTRANSL.langexceldata import Langexceldata, Metainfo,attrjs2key,attrkey2js
from SSOT_db.IM_JSON import JSModel
from SSOT_infra import nvl


class Exportdata:
    def __init__(self, pjsonfile):
        self._model: JSModel = JSModel.readfromfile(pfilename=pjsonfile)

        self._deflang = self._model.jsmodel['model']['language']
        self._languages = self._model.jsmodel['languages'].keys()
        self._metainfo = Metainfo(gitrevision=self._model.jsmodel["_imprint_"]["git-revision"],
                                  dbmodelversion=self._model.jsmodel["_imprint_"]["Modelversion"],
                                  lastupdate=nvl(self._model.jsmodel["model"]["dm"],
                                                 self._model.jsmodel["model"]["dc"]),
                                  jsonfile=pjsonfile,
                                  modellang=self._model.jsmodel["model"]["language"],
                                  modelname=self._model.jsmodel["model"]["name"])
        self.gatherdata()

    def fulldata(self, ptext: dict, pdesc: str, pcomment: str) -> dict:
        retval = ptext
        retval['descr'] = pdesc
        retval['comment'] = pcomment
        return retval

    @staticmethod
    def xlskey(pid: str, pattr: str, pidx: int = None):
        retval = pid + '-' + pattr
        if pidx is not None:
            retval += '-' + str(pidx)
        return retval

    def metainfo(self):
        return self._metainfo

    def getdata(self, ptype=None):
        if ptype is None:
            return self._data
        else:
            return self._data[ptype]

    def _getdeflangstr(self, pelem):
        return pelem[self._deflang]

    def getlanguages(self):
        return self._languages

    def _getdefname(self, pelem):
        return self._getdeflangstr(pelem['name'])

    def _getelement(self, pelemid):
        return self._model.getbyid(pelemid)

    def _readentidata(self):
        elements = self._model.getelements('entities')
        for key, elem in elements.items():
            self._data[self.xlskey(key, 'name')] = self.fulldata(elem[attrkey2js('name')], self._getdefname(elem) + '  Entity-Name',
                                                                 '')
            self._data[self.xlskey(key, 'descr')] = self.fulldata(elem[attrkey2js('descr')],
                                                                  self._getdefname(elem) + '  Entity--Description', '')
            self._data[self.xlskey(key, 'tooltip')] = self.fulldata(elem[attrkey2js('tooltip')],
                                                                    self._getdefname(elem) + '  Entity--Tooltip', '')
            for idx, syno in enumerate(elem[attrkey2js('synonyms')], start=1):
                self._data[self.xlskey(key, 'synonyms', idx)] = self.fulldata(syno, self._getdefname(
                    elem) + '->' + self._getdeflangstr(syno) + '  - Synonym-' + str(idx), '')
            for idx, expl in enumerate(elem[attrkey2js('examples')], start=1):
                self._data[self.xlskey(key, 'examples', idx)] = self.fulldata(expl, self._getdefname(
                    elem) + '  - Example-' + str(idx), '')

    def _readburudata(self):
        elements = self._model.getelements('businessrules')
        for key, elem in elements.items():
            refname = self._getdefname(elem)
            self._data[self.xlskey(key, 'name')] = self.fulldata(elem[attrkey2js('name')], refname + '  Businessrule-Name', '')
            self._data[self.xlskey(key, 'descr')] = self.fulldata(elem[attrkey2js('descr')], refname + '  Businessrule-Description',
                                                                  '')
            self._data[self.xlskey(key, 'errormsg')] = self.fulldata(elem[attrkey2js('errormsg')],
                                                                     refname + '  Businessrule-Errormessage', '')

    def _readattrdata(self):
        elements = self._model.getelements('attributes')
        for key, elem in elements.items():
            refname = self._getdefname(self._getelement(elem['entity'])) + '->' + self._getdefname(elem)
            self._data[self.xlskey(key, 'name')] = self.fulldata(elem[attrkey2js('name')], refname + '  Attribute-Name', '')
            self._data[self.xlskey(key, 'descr')] = self.fulldata(elem[attrkey2js('descr')], refname + '  Attribute-Description',
                                                                  '')
            self._data[self.xlskey(key, 'tooltip')] = self.fulldata(elem[attrkey2js('tooltip')], refname + '  Attribute-Tooltip',
                                                                    '')
            for idx, expl in enumerate(elem[attrkey2js('examples')], start=1):
                self._data[self.xlskey(key, 'examples', idx)] = self.fulldata(expl, self._getdefname(
                    elem) + '  - Example-' + str(idx), '')

    def _readdomadata(self):
        elements = self._model.getelements('domains')
        for key, elem in elements.items():
            refname = self._getdefname(elem)
            self._data[self.xlskey(key, 'name')] = self.fulldata(elem[attrkey2js('name')], refname + '  Domain-Name', '')
            self._data[self.xlskey(key, 'descr')] = self.fulldata(elem[attrkey2js('descr')], refname + '  Domain-Description', '')
        return

    def _readreladata(self):
        elements = self._model.getelements('relations')
        for key, elem in elements.items():
            fromentiname = self._getdefname(self._getelement(elem[attrkey2js('fromto')]['enti']))
            toentiname = self._getdefname(self._getelement(elem[attrkey2js('tofrom')]['enti']))
            self._data[self.xlskey(key, 'fromto')] = self.fulldata(elem[attrkey2js("fromto")]['assoc'],
                                                                   'Relation -' + fromentiname + ' => ' + toentiname,
                                                                   '')
            self._data[self.xlskey(key, 'tofrom')] = self.fulldata(elem[attrkey2js('tofrom')]['assoc'],
                                                                   'Relation -' + toentiname + ' => ' + fromentiname,
                                                                   '')

    def gatherdata(self):
        self._data = {}
        self._readentidata()
        self._readattrdata()
        self._readdomadata()
        self._readreladata()
        self._readburudata()


# Exportdata

def writesheets(pwb, pdata: Exportdata):
    ws = pwb.active
    ws.protection.sheet = True

    excel = Langexceldata()
    langs = pdata.getlanguages()
    excel.setheader(plangs=langs)

    excel.setheaderwidth(pdimensions=[20] + [50] * len(langs) + [45, 40])

    ws.append(excel.getheaderlist())
    ws["A1"].comment = Comment(text=pdata.metainfo(),
                               author="SPOD-Generator", height=100, width=400)

    for idx, d in enumerate(excel.getheaderwidth(), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = d

    for key, data in pdata.getdata().items():
        row = [key]
        for l in langs:
            row.append(data[l])
        row.append(data['descr'])
        row.append(data['comment'])
        ws.append(row)

    # unlock all cells but head and first
    for row in ws.rows:
        for cell in row:
            if cell.row > 1:
                if cell.column > 1:
                    cell.protection = styles.Protection(locked=False)
                cell.alignment = styles.Alignment(wrapText=True, vertical='top')

    return


def createexcel(pdestfile, pjsonfile):
    data = Exportdata(pjsonfile)
    # getdata = {key: {lang:str,} 'desc':desc,'comment':comment}
    wb = Workbook()
    writesheets(pwb=wb, pdata=data)
    wb.save(filename=pdestfile)
    print(f"""Translation excel generated: {pdestfile} """)
    return


def createlangexcel(pjsonfile, pdest=None):
    assert os.path.isfile(pjsonfile)
    destdir = os.path.dirname(pjsonfile)
    destfilename = os.path.basename(pjsonfile)[:-5] + '.xlsx'
    if pdest is None:
        destfile = os.path.join(destdir, destfilename)
    elif os.path.isdir(pdest):
        destfile = os.path.join(pdest, destfilename)
    else:
        # assume it is a full file spec
        destfile = pdest
    createexcel(pdestfile=destfile, pjsonfile=pjsonfile)
    return


if __name__ == '__main__':
    createlangexcel(pjsonfile=sys.argv[1], pdest=None if len(sys.argv) < 3 else sys.argv[2])
