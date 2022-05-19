import os
import tempfile
import unittest

import SSOT_infra.tests.integration as tb
from LOAD_MODELS.LOAD_INFRA import mergedbs
from SSOT_db.IM_JSON import JSModel
from tools.LANGTRANSL import exportdata, importdata, langexceldata


def nocomments(pws):
    for r in pws.iter_rows(min_row=2):
        if r[5].value != '':
            return False
    return True


class MyTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.tm2 = tb.Testmodel(tb.TESTMODEL2)
        self.tm2.initDB(palways=True)

        from openpyxl import Workbook

        self.testdata = [
            ['Key', 'de', 'en', 'fr', 'Description', 'Comments'],
            ['ENTI119-name', 'Hauptentität', 'Master Entity', 'Entité principale', 'Hauptentität  Entity-Name', ''],
            ['ENTI118-name', 'Einzelentität', 'Single Entity', 'Entité unique', 'Hauptentität  Entity-Name', ''],
            ['ENTI119-descr', """Master-Entität mit 3 Kindern mit Attributen und Klassifizierungen
Wird auf allen Zoomstufen angezeigt (0-4)
einzelnes Attribut Eindeutiger Schlüssel""",
             """Master entity with 3 children with attributes and classifications
Displayed on all zoom levels (0-4)
single attribute Unique key""",
             """*Entité maître avec 3 enfants avec attributs et classifications
Affiché sur tous les niveaux de zoom (0-4)
attribut unique Clé unique""", 'Hauptentität  Entity--Description', ''],
            ['ENTI119-tooltip', '', '', '', 'Hauptentität  Entity--Tooltip', ''],
            ['ENTI119-synonyms-1', 'Vaterentität', 'Main Entity', 'Entité maître',
             'Hauptentität->Vaterentität  - Synonym-1', ''],
            ['ENTI119-examples-1', '1. Beispiel Master-Entität', '1. Example: Master entity',
             '1. exemple: Entité maître', 'Hauptentität  - Example-1', ''],
            ['ENTI119-examples-2', '2. Beispiel Master-Entität.', '2. Example: Master Entity',
             '2. exemple: Entité maître', 'Hauptentität  - Example-2', ''],
            ['ATTR113-examples-1', '1. Beispiel Attribute', '1. Beispiel Attribute first appearance',
             '1. Example: Attribut', 'erste Erscheinung  - Example-1', ''],
            ['ATTR123-name', 'Name', 'Name', 'Nom', 'Hauptentität->Name  Attribute-Name', ''],
            ['ATTR123-descr', 'Beschreibendes, obligatorisches Attribut mit Bereich, angezeigt auf Zoomstufe 0-2',
             'Descriptive, mandatory Attribute with domain, showed on zoom level 0-2',
             'Attribut descriptif, obligatoire avec domaine, affiché au niveau de zoom 0-2',
             'Hauptentität->Name  Attribute-Description', ''],
            ['ATTR123-tooltip', '', '', '', 'Hauptentität->Name  Attribute-Tooltip', ''],
            ['DOMA90-name', 'XMLTYPE', 'XMLTYPE', 'XMLTYPE', 'XMLTYPE  Domain-Name', ''],
            ['DOMA90-descr', '', '', '', 'XMLTYPE  Domain-Description', ''],
            ['DOMA73-name', 'email Address', 'email Address', 'email Address', 'email Address  Domain-Name', ''],
            ['DOMA73-descr', 'Email address with regexp check', 'Email address with regexp check',
             'Email address with regexp check', 'email Address  Domain-Description', ''],
            ['RELA138-fromto', '', '', '', 'Relation -Kind Entität1 => Hauptentität', ''],
            ['RELA138-tofrom', '', '', '', 'Relation -Hauptentität => Kind Entität1', ''],
            ['RELA126-fromto', 'fügt hinzu', 'adds', 'ajoute', 'Relation -Zusätzliche Einheit => Rollen-Entität 2', ''],
            ['RELA126-tofrom', 'hinzugefügt von', 'added by', 'ajouté par',
             'Relation -Rollen-Entität 2 => Zusätzliche Einheit', ''],
            ['ATTR113-name', 'erste Erscheinung', 'First Appearance', 'première apparition',
             'Kind Entität1->erste Erscheinung', 'Attribute-Name']
        ]
        self.wb = Workbook()
        self.testws = self.wb.active
        for r in self.testdata:
            self.testws.append(r)
        # destfilename=self.tm2.dbdir/"testexcel.xlsx"
        # wb.save(filename=destfilename)
        self.testexcel = langexceldata.Langexceldata()
        self.testexcel.analyzeheader(self.testws['1'])
        self.testjson = JSModel.readfromfile(
            pfilename=self.tm2.jsonfile)

    def test_metainfo(self):
        self.assertTrue(langexceldata.strislang('xx'))
        self.assertFalse(langexceldata.strislang('x'))
        self.assertFalse(langexceldata.strislang(None))
        self.assertFalse(langexceldata.strislang(''))
        self.assertFalse(langexceldata.strislang('xY'))

        mi = langexceldata.Metainfo(modelname='test')
        print(mi)
        return

    def test_excelcheck(self):

        # all but header is returned as OK rows
        self.assertEqual(len(self.testdata) - 1,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertTrue(nocomments(self.testws))

        # invalid key
        self.testws.cell(row=2, column=1).value = 'gugus'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'ENTI99999-name'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'ENTI99999-name'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'ENTI-synonyms-a'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'ENTI-synonyms-99'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'ENTI-synonyms'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # invalid key
        self.testws.cell(row=2, column=1).value = 'RELA-to-from'
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=1).value = self.testdata[1][0]

        # empty not null column de
        self.testws.cell(row=2, column=2).value = ''
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=2).value = self.testdata[1][1]
        # empty not null column en
        self.testws.cell(row=2, column=3).value = ''
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=2, column=3).value = self.testdata[1][2]

        # duplicate name de
        self.testws.cell(row=3, column=2).value = self.testdata[1][1]
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=3, column=2).value = self.testdata[2][1]

        # duplicate name fr
        self.testws.cell(row=3, column=4).value = self.testdata[1][3]
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=3, column=4).value = self.testdata[2][3]

        # two errors in line
        self.testws.cell(row=3, column=4).value = self.testdata[1][3]
        self.testws.cell(row=3, column=3).value = self.testdata[1][2]
        self.assertEqual(len(self.testdata) - 2,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.wb.save(self.tm2.dbdir / "testexcel.xlsx")
        self.assertEqual(importdata.redfill.fgColor.value, self.testws.cell(row=3, column=3).fill.fgColor.value)
        self.assertEqual(importdata.redfill.fgColor.value, self.testws.cell(row=3, column=4).fill.fgColor.value)
        self.assertFalse(nocomments(self.testws))
        self.testws.cell(row=3, column=4).value = self.testdata[2][3]
        self.testws.cell(row=3, column=3).value = self.testdata[2][2]

        # duplicate attrnames are no error
        self.testws.cell(row=10, column=2).value = self.testws.cell(row=21, column=2).value
        self.assertEqual(len(self.testdata) - 1,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.wb.save(self.tm2.dbdir / "testexcel.xlsx")
        self.assertEqual(importdata.emptyfill.fgColor.value, self.testws.cell(row=21, column=2).fill.fgColor.value)
        self.assertTrue(nocomments(self.testws))
        self.testws.cell(row=10, column=2).value = self.testdata[8][1]

        # all comments are removed if it checks ok
        self.assertEqual(len(self.testdata) - 1,
                         len(importdata.checkentries(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)))
        self.assertTrue(nocomments(self.testws))
        # os.remove(self.tm2.dbdir / "testexcel.xlsx")
        return

    def test_calls(self):
        with self.assertRaises(Exception) as exp:
            exportdata.createlangexcel('gugus')
        print("")

        data = exportdata.Exportdata(pjsonfile=self.tm2.jsonfile).getdata()
        for key in ('ENTI118-name', 'ENTI118-descr', 'ENTI118-tooltip', 'ENTI111-synonyms-1',
                    'ATTR123-name', 'ATTR123-descr', 'ATTR123-tooltip',
                    'DOMA90-name', 'DOMA90-descr', 'ENTI119-examples-2'):
            self.assertTrue(key in data)

        destfilename = os.path.join(self.tm2.dbdir, self.tm2.modelname + '.xlsx')
        if os.path.exists(destfilename):
            os.remove(destfilename)
        exportdata.createlangexcel(self.tm2.jsonfile)
        self.assertTrue(os.path.exists(destfilename))
        os.remove(destfilename)
        exportdata.createlangexcel(self.tm2.jsonfile, pdest=self.tm2.dbdir)
        self.assertTrue(os.path.exists(destfilename))
        os.remove(destfilename)
        exportdata.createlangexcel(self.tm2.jsonfile, pdest=destfilename)
        self.assertTrue(os.path.exists(destfilename))
        # os.remove(destfilename)

        return

    def test_exceljson(self):

        # no changes
        changes, newjson = importdata.mergeexcel2json(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)
        self.assertEqual(0, changes)
        for key, val in newjson.getelements('entities').items():
            self.assertEqual(self.testjson.getbyid(key)['name']['de'], val['name']['de'])
            self.assertEqual(self.testjson.getbyid(key)['name']['fr'], val['name']['fr'])
            self.assertEqual(self.testjson.getbyid(key)['name']['en'], val['name']['en'])
            self.assertEqual(self.testjson.getbyid(key)['tooltip']['de'], val['tooltip']['de'])
        for key, val in newjson.getelements('domains').items():
            self.assertEqual(self.testjson.getbyid(key)['name']['de'], val['name']['de'])
            self.assertEqual(self.testjson.getbyid(key)['name']['fr'], val['name']['fr'])
            self.assertEqual(self.testjson.getbyid(key)['name']['en'], val['name']['en'])
            self.assertEqual(self.testjson.getbyid(key)['descr']['de'], val['descr']['de'])

        # change 1 name
        self.testws.cell(row=2, column=2).value = 'gugus'
        changes, newjson = importdata.mergeexcel2json(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)
        # self.wb.save(self.tm2.dbdir / "testexcel.xlsx")
        self.assertEqual(1, changes)
        self.assertTrue(nocomments(self.testws))
        self.assertEqual(importdata.greenfill.fgColor.value, self.testws.cell(row=2, column=2).fill.fgColor.value)
        self.testws.cell(row=2, column=2).value = self.testdata[1][1]
        self.testws.cell(row=2, column=2).fill = importdata.emptyfill
        self.assertEqual('gugus', newjson.getbyid('ENTI119')['name']['de'])

        changes, newjson = importdata.mergeexcel2json(pws=self.testws, pexcel=self.testexcel, pjson=newjson)
        # self.wb.save(self.tm2.dbdir / "testexcel.xlsx")
        self.assertEqual(1, changes)
        self.assertTrue(nocomments(self.testws))
        self.assertEqual(importdata.greenfill.fgColor.value, self.testws.cell(row=2, column=2).fill.fgColor.value)
        self.assertEqual(self.testdata[1][1], newjson.getbyid('ENTI119')['name']['de'])

        # second change
        self.testws.cell(row=2, column=2).value = 'gugus2'
        self.testws.cell(row=5, column=3).value = 'Tooltip Test'
        changes, newjson = importdata.mergeexcel2json(pws=self.testws, pexcel=self.testexcel, pjson=self.testjson)
        # self.wb.save(self.tm2.dbdir / "testexcel.xlsx")
        self.assertEqual(2, changes)
        self.assertEqual(importdata.greenfill.fgColor.value, self.testws.cell(row=2, column=2).fill.fgColor.value)
        self.assertEqual(importdata.greenfill.fgColor.value, self.testws.cell(row=5, column=3).fill.fgColor.value)
        self.testws.cell(row=2, column=2).value = self.testdata[1][1]
        self.testws.cell(row=5, column=3).value = self.testdata[4][2]
        self.assertEqual('gugus2', newjson.getbyid('ENTI119')['name']['de'])
        self.assertEqual('Tooltip Test', newjson.getbyid('ENTI119')['tooltip']['en'])
        # os.remove(self.tm2.dbdir / "testexcel.xlsx")

    def test_importdata(self):
        from openpyxl import load_workbook
        import shutil
        from SSOT_db.SQL_INFRA import dbConnect
        from SSOT_db.IM_OBJECTS import Entity

        resultjson = str(self.tm2.jsonfile).replace('.json', '_result.json')

        with self.assertRaises(Exception) as exp:
            importdata.importlangexcel('gugus')
        print("")
        impfilename = os.path.join(self.tm2.dbdir, self.tm2.modelname + '.xlsx')
        importdata.importlangexcel(impfilename)
        self.assertFalse(os.path.exists(resultjson))  # nothing changed, no files generated

        with tempfile.TemporaryDirectory() as tempdir:
            os.chdir(tempdir)
            shutil.copyfile(impfilename, 'myinput.xlsx')
            shutil.copyfile(self.tm2.jsonfile, 'myjson.json')
            shutil.copyfile(self.tm2.dbfile, 'testmodel-2.db')
            wb = load_workbook(filename='myinput.xlsx')
            ws = wb.active
            a1comment = ws["A1"].comment
            self.assertIsNotNone(a1comment)
            a1comment.text = a1comment.text.replace(str(self.tm2.jsonfile), 'myjson.json')
            excel = langexceldata.Langexceldata()
            excel.analyzecomment('' if a1comment is None else a1comment.text)
            self.assertEqual('myjson.json', excel.getmetainfo().getjsonfile())
            ws["A1"].comment = a1comment
            wb.save(filename="myinput.xlsx")
            # wb.save(filename=impfilename.replace('.xlsx','_test.xlsx'))
            changes = importdata.importlangexcel("myinput.xlsx")
            self.assertEqual(0, changes)
            # shutil.copyfile("myinput_result.xlsx",self.tm2.dbdir/"myinput_result.xlsx")
            self.assertFalse(os.path.exists("myjson_result.json"))
            self.assertFalse(os.path.exists("myinput_result.xlsx"))

            # teste Änderung im Excel
            cell = ws["B2"]
            cell.value = cell.value + "XX"
            cell = ws["C2"]
            cell.value = cell.value + "YY"
            wb.save(filename="myinput.xlsx")
            wb.save(filename=impfilename.replace('.xlsx', '_test.xlsx'))
            changes = importdata.importlangexcel("myinput.xlsx")
            self.assertTrue(changes > 0)

            resultjson = JSModel.readfromfile(pfilename="myjson_result.json")
            mergedbs.mergejs2db(pmodel=resultjson, psrcname="TRANSL", pverbose=False, pdbfile="testmodel-2.db")
            self.assertTrue(os.path.exists("myjson_result.json"))
            self.assertTrue(os.path.exists("myinput_result.xlsx"))

            dbConnect.openDB(pfilepath="testmodel-2.db")
            enti = Entity().getbyid(118)
            self.assertTrue(enti.enti_name.endswith("XX"))
            self.assertTrue(enti.enti_name_l['en'].endswith("YY"))

            dbConnect.closeDB()

            shutil.copyfile("myinput_result.xlsx", self.tm2.dbdir / "myinput_result.xlsx")
            shutil.copyfile("myinput.xlsx", self.tm2.dbdir / "myinput.xlsx")
            shutil.copyfile("myjson_result.json", self.tm2.dbdir / "myjson_result.json")

            # self.assertTrue(os.path.exists(resultjson))

    def test_translateexcel(self):
        with self.assertRaises(Exception) as exp:
            importdata.translateexcel(None, None)

        if os.path.exists("/Users/stb/.deepl/deeplauthid"):
            with open("/Users/stb/.deepl/deeplauthid") as d:
                deeplid = d.read()
        else:
            print("******* Test with real deeplid skipped")
            return

        with self.assertRaises(Exception) as exp:
            importdata.translateexcel(None, deeplid)

        with tempfile.TemporaryDirectory() as tempdir:
            self.wb.save("testexcel.xlsx")
            importdata.translateexcel("testexcel.xlsx", deeplid)

        return


if __name__ == '__main__':
    unittest.main()
