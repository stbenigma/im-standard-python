import sys,re

import openpyxl
def getcolumns (pinwb):
    ws=pinwb["Features"]
    cols = {}
    for rownum,row in enumerate(ws):
        if rownum < 3:continue
        cols[row[1].value] = [v.value for v in row]
    return cols

def transform(poutws,pinwb):
    columns = getcolumns(pinwb)

    poutws.cell(row=1,column=1).value = "PT-SIMUS"
    poutws.cell(row=1,column=3).value = "Information Model"
    poutws.cell(row=2,column=1).value = "tableName"
    poutws.cell(row=2,column=2).value = "columnName"
    poutws.cell(row=2,column=3).value = "entityName"
    poutws.cell(row=2,column=4).value = "attrName"
    currow = 2
    for ws in pinwb.worksheets:
        if re.match("X1F_SIM_C\d{6}",ws.title):
            tablename=ws["B4"].value
            currow += 1
            poutws.cell(row=currow, column=1).value = tablename
            #print (ws.title,ws["B4"].value)
            for col in range(9,500):
                colname = ws[f"B{str(col)}"].value
                if colname is not None:
                    currow += 1
                    poutws.cell(row=currow, column=1).value = tablename
                    poutws.cell(row=currow, column=2).value = colname
                    poutws.cell(row=currow, column=5).value = ws[f"C{str(col)}"].value
                    poutws.cell(row=currow, column=6).value = columns[ws[f"C{str(col)}"].value][4]
                    poutws.cell(row=currow, column=7).value = columns[ws[f"C{str(col)}"].value][5]
                    poutws.cell(row=currow, column=8).value = columns[ws[f"C{str(col)}"].value][7]
                    poutws.cell(row=currow, column=9).value = columns[ws[f"C{str(col)}"].value][11]
                    poutws.cell(row=currow, column=10).value = columns[ws[f"C{str(col)}"].value][12]
                    #print(tablename,colname)

    return

def createExcel(poutfile: str, pinwb):
    outwb = openpyxl.Workbook()
    ws = outwb.create_sheet("PT-SIMUS")
    transform(poutws=ws, pinwb=pinwb)
    outwb.remove(outwb.worksheets[0])
    outwb.save(filename=poutfile)
    return

def main(pinfile,poutfile):
    createExcel(pinwb=openpyxl.load_workbook(filename=pinfile),poutfile=poutfile)
    return


if __name__ == '__main__':
    main(pinfile=sys.argv[1],poutfile=sys.argv[2])
