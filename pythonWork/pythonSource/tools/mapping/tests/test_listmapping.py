import os.path
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from SSOT_db.IM_JSON import JSModel
from SSOT_infra.tests import integration as tb
from tools.mapping import listmapping


class ListMapping(unittest.TestCase):
    def setUp(self) -> None:
        self.crm = tb.ModelHelper(tb.CRMTEST)
        self.crm.initDB(palways=False)

    def test_createAllMapping(self):
        with tempfile.TemporaryDirectory() as tempdir:
            testfile = Path(tempdir + "/testfile.json")
            shutil.copy(self.crm.jsonfile, testfile)
            mapfile = listmapping.createAllMapping(pjsonfile=testfile,
                                                   pdestination=testfile.with_suffix('.xlsx'))
            wb = load_workbook(filename=mapfile)

            # DEBUG
            #shutil.copy(mapfile,Path.home()/"Downloads")
            #print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

            self.assertEqual("Overview", wb.sheetnames[0])
            wssf = wb.get_sheet_by_name("CRM-Salesforce")
            self.assertEqual("person", wssf.cell(153, 1).value)
            self.assertEqual("Geschlecht", wssf.cell(153, 3).value)
            self.assertEqual("natürliche Person", wssf.cell(153, 12).value)
            self.assertEqual("Geschlecht", wssf.cell(153, 13).value)
            self.assertEqual("natural_person", wssf.cell(153, 18).value)
            self.assertEqual("Gender", wssf.cell(153, 19).value)

            locmodel = JSModel.readfromfile(self.crm.jsonfile)
            tb.crmmapentityhack(locmodel)
            mapfile = Path(tempdir + '/map.xlsx')
            listmapping.writeintfxls(pfilename=str(mapfile), pmodel=locmodel,
                                     plang=locmodel.jsmodel["model"]["language"])
            wb = load_workbook(filename=mapfile)
            # shutil.copy(mapfile,Path.home()/"Downloads")
            # print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

            wssf = wb.get_sheet_by_name('Overview')
            self.assertEqual(locmodel.jsmodel["model"]["language"], wssf.cell(16, 2).value)

            wssf = wb.get_sheet_by_name("CRM-Salesforce")
            self.assertEqual("GeoInformation", wssf.cell(43, 3).value)
            self.assertEqual("col1tomany", wssf.cell(43, 27).value)

            wssf = wb.get_sheet_by_name("TestMapping")
            self.assertEqual("col1tomany", wssf.cell(4, 3).value)
            self.assertTrue(wssf.cell(4, 12).value.endswith("Geografische Einheit (Administrativgebiet)"))
            self.assertTrue(wssf.cell(4, 13).value.endswith("GeoInformation"))

            listmapping.writeintfxls(pfilename=str(mapfile), pmodel=locmodel,
                                     plang="en")
            wb = load_workbook(filename=mapfile)

            wssf = wb.get_sheet_by_name('Overview')
            self.assertEqual("en", wssf.cell(16, 2).value)

            wssf = wb.get_sheet_by_name("TestMapping")
            self.assertTrue( wssf.cell(4, 12).value.endswith("Geografical Unit (Administrative Territory)"))
            # shutil.copy(mapfile,Path.home()/"Downloads")
            # print (f"DEBUG: copied xls to {Path.home()/'Downloads'}")

        return

    def test_specialjson(self):
        return  # for local customer-jsons
        with tempfile.TemporaryDirectory() as tempdir:
            testfile = tempdir + "/testfile.json"
            shutil.copy('/Users/stb/Documents/Projekte/SPODtools/testdata/localtestmodels/.......json', testfile)
            mapfile = listmapping.createAllMapping(pjsonfile=testfile)
            self.assertTrue(os.path.exists((tempdir + "/testfile.json")))
            return

    def test_call(self):
        sys.argv = [
            listmapping.__file__,
            '--language', 'en',
            str(self.crm.jsonfile)
        ]
        listmapping.main()
        sys.argv = [
            listmapping.__file__,
            str(self.crm.jsonfile)
        ]
        with tempfile.TemporaryDirectory() as tempdir:
            testfile = tempdir + "testexcel.xlsx"
            listmapping.main()
            sys.argv = [
                listmapping.__file__,
                '--destination', testfile,
                str(self.crm.jsonfile)
            ]
            listmapping.main()
            self.assertTrue(os.path.isfile(testfile))

        return


if __name__ == '__main__':
    unittest.main()
