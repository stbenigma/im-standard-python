import sys
from openpyxl import Workbook
from SSOT_db.IM_JSON import  JSModel

# utils.get_column_letter(pidx)
from tools.mapping import writeoverview

val2str = lambda v: '' if v is None else str(v) if type(v) in (int, float) else v


def writesheets(pwb, pjsmodel: JSModel):
    systems = pjsmodel.getelements('INTF')
    cols = pjsmodel.getelements('COLU')
    writeoverview(pwb=pwb,pmodel=pjsmodel)

    for sys in systems.values():
        ws = pwb.create_sheet(sys["name"])
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 70
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 60
        tables = [pjsmodel.getbyid(tab) for tab in sys["tables+"]]
        tables = sorted(tables, key=lambda tab: tab["name"])
        rowidx, colidx = 1, 1
        ws.cell(column=colidx, row=rowidx, value='Table')
        ws.cell(column=colidx + 1, row=rowidx, value='Column')
        ws.cell(column=colidx + 2, row=rowidx, value='Entity/Relation')
        ws.cell(column=colidx + 3, row=rowidx, value='Attribute')
        rowidx += 1
        for tab in tables:
            colidx = 1
            if len(tab['entitiesmapped']) + len(tab['relationsmapped']) == 0:
                # no entity write only table and columns
                ws.cell(column=colidx, row=rowidx, value=tab["name"])
                rowidx += 1
            else:
                for map in tab['entitiesmapped']:
                    ws.cell(column=colidx, row=rowidx, value=tab["name"])
                    ws.cell(column=colidx + 2, row=rowidx, value=pjsmodel.getbyid(map)["name"]["de"])
                    rowidx += 1
                for map in tab['relationsmapped']:
                    ws.cell(column=colidx, row=rowidx, value=tab["name"])
                    ws.cell(column=colidx + 2, row=rowidx, value=pjsmodel.getbyid(map)["name"])
                    rowidx += 1
            # fi

            columns = [pjsmodel.getbyid(col) for col in tab["columns+"]]
            columns = sorted(columns, key=lambda col: col["name"])
            for col in columns:
                colidx = 1
                if len(col["attributesmapped"]) == 0:
                    ws.cell(column=colidx + 1, row=rowidx, value=col["name"])
                    rowidx += 1
                else:
                    for attrid in col["attributesmapped"]:
                        attr = pjsmodel.getbyid(attrid)
                        ws.cell(column=colidx + 1, row=rowidx, value=col["name"])
                        ws.cell(column=colidx + 2, row=rowidx, value=pjsmodel.getbyid(attr["entity"])["name"]["de"])
                        ws.cell(column=colidx + 3, row=rowidx, value=attr["name"]["de"])
                        rowidx += 1
                #fi
            # for
        # for
    # for
    return


def createExcel(pfilename: str, pjsmodel):
    wb = Workbook()
    writesheets(pwb=wb, pjsmodel=pjsmodel)
    wb.remove(wb.worksheets[0])
    wb.save(filename=pfilename)
    return


def createMapExcel(pjsonfile):
    JSONEXTENSION = '.json'
    jsmodel = JSModel.readfromfile(pfilename=pjsonfile)
    filename = pjsonfile[:-len(JSONEXTENSION)] + '_MAP'
    fileext = '.xlsx'
    createExcel(pfilename=filename + fileext, pjsmodel=jsmodel)
    print("Excel {} created".format(filename + fileext))


if __name__ == '__main__':
    createMapExcel(pjsonfile=None if len(sys.argv) == 1 else sys.argv[1])
