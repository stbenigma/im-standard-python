# -*- coding: latin-1 -*-
import os
import re
import sys

sys.path.append(os.path.dirname(os.path.realpath(__file__)) + '/../IM_db')
sys.path.append(os.path.dirname(os.path.realpath(__file__)) + '/..')
from IM_DB import *
from IM_OBJECTS import *
import __main__
from openpyxl import Workbook
from openpyxl.styles import Alignment

fileCSV = None
EOL: str = '\n'
CSVSEP: str = ';'
entities = {}
attributes = {}
tables = {}
columns = {}
schnittstellen = {}
tabentimap = {}
colattrmap = {}


def colnum_string(n):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string
# colum_string

def cell_string(n,m):
    string = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string+str(m)
#cell_string

def setcell(pws,pcolumn, prow, pvalue
            ,phorizontal = None, pvertical=None
            ,ptext_rotation=None, pwrap_text = None):

    cell = pws.cell(column=pcolumn, row=prow, value=pvalue)
    if (phorizontal is not None or pvertical is not None\
        or ptext_rotation is not None or pwrap_text is not None):
        cell.alignment = Alignment(horizontal=phorizontal
                                       , vertical=pvertical
                                       , text_rotation=ptext_rotation
                                        ,wrap_text = pwrap_text
                                       )
    #if
#setcell

def writesheettabent(pwb: Workbook):
    ws = pwb.create_sheet("Table to Entity mapping")
    rowidx, colidx = 1, 1
    ws.cell(column=colidx, row=rowidx, value='Interface')
    colidx += 1
    ws.cell(column=colidx, row=rowidx, value='Table')
    colidx += 1
    for idx, enti in enumerate(entities.values()):
        setcell(pws=ws,pcolumn=idx + colidx, prow=rowidx, pvalue=enti[0]
                ,ptext_rotation=90)
    # for
    rowidx += 1
    for skey, sval in schnittstellen.items():
        for tkey, tval in sval.items():
            colidx = 1
            setcell(pws=ws,pcolumn=colidx, prow=rowidx, pvalue=skey
                    )
            colidx += 1
            ws.cell(column=colidx, row=rowidx, value=tval[0])
            colidx += 1
            for entiid in entities.keys():
                #            for idx, tabkey in enumerate(tables[tkey][3]):
                #                if (tabkey == 'X'):
                if (istintabentimap(tkey, entiid)):
                    setcell (pws= ws,pcolumn=colidx, prow=rowidx, pvalue='X'
                            ,phorizontal="center", pvertical="center")
                # fi
                colidx += 1
            # for
            rowidx += 1
        # for
    # for

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    for idx in range(3, colidx): ws.column_dimensions[colnum_string(idx)].width = 3
# writesheettabent

def writesheetentitab(pwb: Workbook):
    ws = pwb.create_sheet("Entity to Table mapping")

    rowidx, colidx = 1, 1
    ws.cell(column=colidx, row=rowidx, value='Information Model')
    ws.cell(column=colidx, row=rowidx + 1, value='Entity')
    colidx += 1
    for schn_name, tabs in schnittstellen.items():
        ws.cell(column=colidx, row=rowidx, value=schn_name)
        for tabkey in tabs.values():
            setcell(pws= ws,pcolumn=colidx, prow=rowidx + 1, pvalue=tabkey[0]
                    , ptext_rotation=90)
            colidx += 1
        # for
    # for
    rowidx += 2

    for entiid, enti in entities.items():
        colidx = 1
        setcell (pws=ws,pcolumn=colidx, prow=rowidx, pvalue=enti[0])
        colidx += 1

        for tabs in schnittstellen.values():
            for tabkey in tabs.keys():
                if (istintabentimap(tabkey, entiid)):
                    setcell(pws = ws,pcolumn=colidx, prow=rowidx + 1, pvalue='X')
                # if
                colidx += 1
            # for
        # for
        rowidx += 1
    # for

    ws.column_dimensions['A'].width = 35
    for idx in range(2, colidx): ws.column_dimensions[colnum_string(idx)].width = 3
# writesheetentitab

def writesheetcolattr(pwb: Workbook):
    ws = pwb.create_sheet("Columns to Attributes mapping")

    rowidx, colidx = 1, 1
    ws.cell(column=colidx + 2, row=rowidx, value='Entity')
    ws.cell(column=colidx, row=rowidx + 1, value='Interface')
    ws.cell(column=colidx + 1, row=rowidx + 1, value='Table')
    ws.cell(column=colidx + 2, row=rowidx + 1, value='Column')
    colidx += 3
    for enti in entities.values():
        ws.cell(column=colidx, row=rowidx, value=enti[0])
        for attr in enti[2].values():
            cell = ws.cell(column=colidx, row=rowidx + 1, value=attr[0])
            cell.alignment = Alignment(horizontal='general'
                                       , vertical='bottom'
                                       , text_rotation=90
                                       # ,wrap_text = False
                                       )
            colidx += 1
        # for
    # for
    rowidx += 2

    for skey, sval in schnittstellen.items():
        for tkey, tval in sval.items():
            colidx = 1
            for ckey, cval in tval[1].items():
                ws.cell(column=colidx, row=rowidx, value=skey)
                colidx += 1
                ws.cell(column=colidx, row=rowidx, value=tval[0])
                colidx += 1
                ws.cell(column=colidx, row=rowidx, value=cval[0])
                colidx += 1
                for enti in entities.values():
                    for attrid in enti[2].keys():
                        if (istincolattrmap(ckey,attrid)):
                            ws.cell(column=colidx, row=rowidx, value='X')
                        # if
                        colidx += 1
                    # for
                # for
                rowidx += 1
            # for
        # for
    # for

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 35
    for idx in range(4, colidx): ws.column_dimensions[colnum_string(idx)].width = 3
# writesheetcolattr

def writesheetattrcol(pwb: Workbook):
    ws = pwb.create_sheet("Attributes to Columns mapping")

    rowidx, colidx = 1, 1
    ws.cell(column=colidx, row=rowidx, value='Information Model')
    ws.cell(column=colidx, row=rowidx + 1, value='Entity')
    colidx += 1
    ws.cell(column=colidx+1, row=rowidx + 1, value='Attribute')
    colidx += 1
    for schn_name, tabs in schnittstellen.items():
        ws.cell(column=colidx, row=rowidx, value=schn_name)
        for tabkey in tabs.values():
            setcell (pws= ws,pcolumn=colidx, prow=rowidx + 1, pvalue=tabkey[0]
                    , ptext_rotation=90 )
            colidx += 1
        # for
    # for
    rowidx += 2

    for enti in entities.values():
        for attrid, attr in enti[2].items():
            colidx = 1
            setcell (pws= ws,pcolumn=colidx, prow=rowidx, pvalue=enti[0]
                    ,pwrap_text = True )
            colidx += 1
            setcell (pws=ws,pcolumn=colidx, prow=rowidx, pvalue=attr[0]
                     ,pwrap_text = True )
            colidx += 1
            for skey, sval in schnittstellen.items():
                for tkey, tval in sval.items():
                    value = ''
                    for ckey, cval in tval[1].items():
                        if (istincolattrmap(ckey,attrid)):
                            value += '' if (value =='')  else EOL
                            value += cval[0]
                        #if
                    #for
                    if (value != ''):
                        setcell(pws=ws, pcolumn=colidx, prow=rowidx, pvalue=value
                            , pwrap_text=True)
                        headcell = ws.cell(column=colidx, row=2)
                        headcell.alignment = Alignment(text_rotation=0, wrap_text=True)
                    colidx += 1
                #for
            #for
            rowidx += 1
        #for
    #for

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    for idx in range(3, colidx):
        width = 5 if (ws.cell(column=idx,row=2).alignment.text_rotation == 90) else 25
        ws.column_dimensions[colnum_string(idx)].width = width
    #for
# writesheetattrcol

def writesheetschnittstelle(pwb, pschnname, pschn):
    ws = pwb.create_sheet(pschnname)

    rowidx, colidx = 1, 1
    ws.cell(column=colidx , row=rowidx, value=pschnname)
    ws.cell(column=colidx + 2 , row=rowidx, value='Information Model')
    ws.cell(column=colidx , row=rowidx + 1, value='tableName')
    ws.cell(column=colidx + 1, row=rowidx + 1, value='columnName')
    ws.cell(column=colidx + 2, row=rowidx + 1, value='entityName')
    ws.cell(column=colidx + 3, row=rowidx + 1, value='attrName')
    rowidx += 2

    for tabid,tab in pschn.items():
        firstrowidx = rowidx
        for entiid,enti in entities.items():
            if (istintabentimap(tabid,entiid)):
                ws.cell(column=colidx, row=rowidx, value=tab[0])
                ws.cell(column=colidx+2, row=rowidx, value=enti[0])
                rowidx += 1
            #if
        #for
        if (firstrowidx == rowidx):
            """keinen Eintrag für eine Entity geschrieben, schreibe die Tabelle sowieso"""
            ws.cell(column=colidx, row=rowidx, value=tab[0])
            rowidx += 1
        #if

        for colid,col in tab[1].items():
            firstrowidx = rowidx
            for enti in entities.values():
                for attrid,attr in enti[2].items():
                    if (istincolattrmap(colid,attrid)):
                        ws.cell(column=colidx, row=rowidx, value=tab[0])
                        ws.cell(column=colidx + 1, row=rowidx, value=col[0])
                        ws.cell(column=colidx + 2, row=rowidx, value=enti[0])
                        ws.cell(column=colidx + 3, row=rowidx, value=attr[0])
                        rowidx += 1
                    #if
                #for
            #for

            if (firstrowidx == rowidx):
                """keinen Eintrag für eine Entity geschrieben, schreibe die Tabelle sowieso"""
                ws.cell(column=colidx, row=rowidx, value=tab[0])
                ws.cell(column=colidx + 1, row=rowidx, value=col[0])
                rowidx += 1
            #if
        #for
    #for
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 35
# writesheetschnittstelle

def writexls(pfilename: str):
    global entities, schnittstellen
    wb = Workbook()

    writesheettabent(pwb=wb)
    writesheetentitab(pwb=wb)
    writesheetattrcol(pwb=wb)
    #    writesheetcolattr(pwb=wb)
    for schnname, schn in schnittstellen.items():
        writesheetschnittstelle(pwb=wb, pschnname=schnname, pschn=schn)
    wb.remove(wb.worksheets[0])

    wb.save(filename=pfilename)


# writexls

def createFile(pfilename):
    global fileCSV
    csvfile = parameters.webDirec() + pfilename
    if os.path.exists(csvfile):
        os.remove(csvfile)
    fileCSV = open(csvfile, 'w')


# createFile

def closefile():
    global fileCSV
    fileCSV.close()


# closefile

def write(*args, **kwargs):
    global fileCSV
    sep = kwargs['sep'] if 'sep' in kwargs else ''
    str = sep.join(arg for arg in args)
    fileCSV.write(str)


def writeln(*args, **kwargs):
    write(*args, **kwargs)
    write(EOL)


def listtabenti():
    global entities
    global attributes
    global tables
    global schnittstellen
    global tabentimap

    createFile(pfilename=parameters.odmModelName() + '_tabenti.csv')
    write('\ufeff')
    topheader = 'Interface' + CSVSEP + 'Table' + CSVSEP + CSVSEP.join(e[0] for e in entities.values())
    writeln(topheader)
    for skey, sval in schnittstellen.items():
        for tkey, tval in sval.items():
            writeln(skey, tval[0], CSVSEP.join(tables[tkey][3]), sep=CSVSEP)
    print("Erstellt: {}".format(fileCSV.name))
    closefile()


# listtabenti


def listcolattr():
    global entities
    global attributes
    global tables
    global schnittstellen
    global tabentimap

    createFile(pfilename=parameters.odmModelName() + '_colattr.csv')
    write('\ufeff')
    topheader = CSVSEP + CSVSEP + 'Entity' + CSVSEP
    subheader = 'Interface' + CSVSEP + 'Table' + CSVSEP + 'Column' + CSVSEP
    for enti in entities.values():
        attrs = enti[2]
        topheader += enti[0] + CSVSEP + CSVSEP.join('' for at in attrs)[:-1]
        subheader += CSVSEP + CSVSEP.join(at[0] for at in attrs.values())
    writeln(topheader)
    writeln(subheader)
    for skey, sval in schnittstellen.items():
        for tkey, tval in sval.items():
            for ckey, cval in tval[1].items():
                writeln(skey, tval[0], cval[0], CSVSEP.join(columns[ckey][3]), sep=CSVSEP)
    print("Erstellt: {}".format(fileCSV.name))
    closefile()
# listcolattr

def istintabentimap(tabid, entiid):
    return (tabid in tabentimap) and (entiid in tabentimap[tabid])
# istintabentimap

def istincolattrmap(colid, attrid):
    return (colid in colattrmap) and (attrid in colattrmap[colid])
# istincolattrmap

def stripeol(str):
    retval = re.sub("\n+", "\n", str)
    retval = retval.strip(EOL)
    return retval
# stripeol

def listentiintf():
    global tabentimap, schnittstellen, entities
    createFile(pfilename=parameters.odmModelName() + '_entiintf.csv')
    write('\ufeff')
    writeln('Information Model', CSVSEP.join(schn_name for schn_name in schnittstellen.keys()), sep=CSVSEP)
    for entiid, enti in entities.items():
        write(enti[0], CSVSEP)
        for schntabs in schnittstellen.values():
            entry = EOL.join(tab[0] if (istintabentimap(tabid, entiid)) else "" for tabid, tab in schntabs.items())
            entry = stripeol(entry)
            if (entry != ""):
                write('"' + entry + '"')
            write(CSVSEP)
        writeln()
    print("Erstellt: {}".format(fileCSV.name))
    closefile()
# listentiintf

def listattrintf():
    global tabentimap, schnittstellen, entities
    createFile(pfilename=parameters.odmModelName() + '_attrintf.csv')
    write('\ufeff')
    topheader = CSVSEP + CSVSEP
    subheader = 'Entity' + CSVSEP + 'Attribute'
    for schn_name, tabs in schnittstellen.items():
        topheader += schn_name + CSVSEP + CSVSEP.join('' for ta in tabs)[:-1]
        subheader += CSVSEP + CSVSEP.join(ta[0] for ta in tabs.values())
    writeln(topheader)
    writeln(subheader)
    for enti_id, enti in entities.items():
        for attrid, attr in enti[2].items():
            write(enti[0], CSVSEP, attr[0], CSVSEP)
            for schntabs in schnittstellen.values():
                entry = []
                for tab in schntabs.values():
                    colentry = EOL.join(
                        (tab[0] + '.' + col[0]) if (istincolattrmap(colid, attrid)) else "" for colid, col in
                        tab[1].items())
                    colentry = stripeol(colentry)
                    if (colentry != ""):
                        entry.append(colentry)
                # for
                entry = EOL.join(entry)
                entry = stripeol(entry)
            if (entry != ""):
                write('"' + entry + '"')
            write(CSVSEP)
            # for
        # for
        writeln()
    # for

    print("Erstellt: {}".format(fileCSV.name))
    closefile()
# listattrintf

def listentitable():
    global entities
    global schnittstellen

    createFile(pfilename=parameters.odmModelName() + '_entitable.csv')
    write('\ufeff')
    topheader = 'Information Model' + CSVSEP
    subheader = 'Entity'
    for schn_name, tabs in schnittstellen.items():
        topheader += schn_name + CSVSEP + CSVSEP.join('' for ta in tabs)[:-1]
        subheader += CSVSEP + CSVSEP.join(ta[0] for ta in tabs.values())
    writeln(topheader)
    writeln(subheader)
    for enti in entities.values():
        writeln(enti[0], CSVSEP.join(enti[3]), sep=CSVSEP)
    print("Erstellt: {}".format(fileCSV.name))
    closefile()
# listentitable

def listtabenti():
    global entities
    global schnittstellen

    createFile(pfilename=parameters.odmModelName() + '_tabenti.csv')
    write('\ufeff')
    topheader = 'Interface' + CSVSEP + 'Table' + CSVSEP + CSVSEP.join(enti[0] for enti in entities.values())
    writeln(topheader)
    for skey, sval in schnittstellen.items():
        for tkey, tval in sval.items():
            writeln(skey, tval[0], CSVSEP.join(tables[tkey][3]), sep=CSVSEP)
    print("Erstellt: {}".format(fileCSV.name))
    closefile()


# listtabenti

def filllists(plang):
    global entities
    global attributes
    global tables
    global columns
    global schnittstellen
    global tabentimap
    global colattrmap
    entities = {enti.enti_id: [enti.enti_name
        , {tem[0]: [t for t in tem[1].keys()]
           for tem in TablEntiMap.tablelist(pentiid=enti.enti_id)}
        , {attr.attr_id: [attr.attr_anzname, attr.attr_tech_name]
           for attr in enti.getattributes()}
                               ]
                for enti in Entitaet.select()}
    attributes = {attr.attr_id: [attr.attr_anzname, attr.attr_tech_name, attr.attr_enti_id, attr.attr_bezi_id] for attr
                  in Attribut.select()}
    tables = {tabl.tabl_id: [tabl.tabl_name
        , Schnittstelle.getname(tabl.tabl_schn_id)
        , {c.scha_id: c.scha_column_name for c in tabl.getcolumns()}
                             ]
              for tabl in Tabelle.select()}

    columns = {scha.scha_id: [scha.scha_column_name, scha.scha_tabl_id, scha.scha_fremdsystem_id] for scha in
               Schnittstelleattr.select()}
    schnittstellen = {schn.schn_name:
                          {tabl.tabl_id: [tabl.tabl_name
                              , {c.scha_id: [c.scha_column_name, c.scha_fremdsystem_id] for c in tabl.getcolumns()}
                                          ] for tabl in Tabelle.selectbyschnid(schn.schn_id)
                           }
                      for schn in Schnittstelle.select()}
    tabentimap = TablEntiMap.extendedtabentimap()
    for tkey, tval in tables.items():
        matentry = lambda tabid, entiid: 'X' if (tabid in tabentimap) and (entiid in tabentimap[tabid]) else ''
        maps = [matentry(tkey, e) for e in entities.keys()]
        tval.append(maps)
    # for
    for ekey, eval in entities.items():
        matentry = lambda tabid, entiid: 'X' if (tabid in tabentimap) and (entiid in tabentimap[tabid]) else ''
        maps = [matentry(tkey, ekey) for tkey in tables.keys()]
        eval.append(maps)
    # for
    colattrmap = AttrTransf.colattrmap()
    for colid, cval in columns.items():
        matentry = lambda colid, attrid: 'X' if (istincolattrmap(colid,attrid)) else ''
        maps = [matentry(colid, attrid) for attrid in attributes.keys()]
        cval.append(maps)
    # for


#    print (len(entities),entities)
#    print(len(attributes),attributes)
#    print (len(tables),tables)
#    print (len(schnittstellen),schnittstellen)
#    print (len(mapping),mapping)
#    print(len(colattrmap),colattrmap)
# filllists

def main(pdirec, plang):
    parameters.initparam(p_callarg=pdirec)
    logging.initlog()

    print("listmapping", parameters.odmBaseDirec(), parameters.odmModelName())

    dbConnect.openDB(p_filepath=parameters.dbFilePath());
    deflang = Sprache.liesdeflangiso2()
    if deflang is not None: parameters.dbDefaultLang(deflang)
    filllists(plang=plang)
    dbConnect.myDbConn.close()
    # listentiintf()
    # listattrintf()
    # listentitable()
    # listtabenti()
    # listcolattr()
    writexls(pfilename=parameters.webDirec() + 'Mappingtables_' + parameters.odmModelName() + '.xlsx')
    logging.logmessage("Mpdel {}: mappinglist form database {}\n  => created in file {}"
          .format(parameters.odmModelName(),parameters.dbFilePath()
                  , parameters.webDirec() + 'Mappingtables_' + parameters.odmModelName() + '.xlsx'))


# main

if __name__ == '__main__':
    direc = sys.argv[1]
    lang = sys.argv[2] if (len(sys.argv) > 2) else None
    main(pdirec=direc, plang=lang)
